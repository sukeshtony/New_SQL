"""
generate_update_sql.py
────────────────────────────────────────────────────────────────
Reads an update Excel template and generates SQL UPDATE queries.
Only the cells you fill in are updated — blank cells = leave as-is.

Usage:
    python generate_update_sql.py                     # uses update_template.xlsx
    python generate_update_sql.py my_updates.xlsx     # custom input file
    python generate_update_sql.py --create-template   # creates a blank template

Output: update_output.sql
"""

import sys
from pathlib import Path
from collections import OrderedDict

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl not installed.  Run: pip install openpyxl")

# ─── Fixed filter constants ────────────────────────────────────────────────────
CREATED_BY = "ee42009d-d83e-4c12-abd1-2d8fff809b18"
ROLE       = "TEACHER"
DIFFICULTY = "EASY"

DEFAULT_EXCEL  = "update_template.xlsx"
OUTPUT_FILE    = "update_output.sql"

# ─── Module definitions ────────────────────────────────────────────────────────
# sheet_name → ordered list of (content_key, json_array_index)
# json_array_index = position of this key inside content_json["contents"]
MODULE_KEYS: dict[str, list[tuple[str, int]]] = {
    "M1 - Reading Paragraph": [
        ("passage",      0),
        ("example",      1),
    ],
    "M2 - Story Building": [
        ("example_question", 0),
        ("example_content",  1),
        ("keywords",         2),
    ],
    "M3 - QA Bursts": [
        ("theme",           0),
        ("rapid_questions", 1),
        ("example_heading", 2),
        ("example_content", 3),
    ],
    "M4 - Presentation": [
        ("topic_overview",         0),
        ("preparation_guidelines", 1),
        ("example_heading",        2),
        ("example_content",        3),
    ],
    "M5 - Tongue Twisters": [
        ("example",       0),
        ("tongue_twister", 1),
    ],
    "M6 - Poems": [
        ("example",       0),
        ("assigned_text", 1),
    ],
    "M7 - Silly Debate": [
        ("topic",            0),
        ("user_side",        1),
        ("ai_side",          2),
        ("content",          3),
        ("example_topic",    4),
        ("example_user_side",5),
        ("example_ai_side",  6),
    ],
    "M8 - Audio Postcards": [
        ("topic",               0),
        ("instructions",        1),
        ("example_topic",       2),
        ("example_instruction", 3),
        ("example_content",     4),
    ],
}

# ─── Helpers ───────────────────────────────────────────────────────────────────
def sq(val: str) -> str:
    """Escape single quotes for PostgreSQL string literals."""
    return str(val).replace("'", "''")


def read_sheet(ws) -> dict[str, list[tuple[str, str]]]:
    """
    Read a module sheet.
    Row 1 = headers (lesson_name, key1, key2, ...)
    Row 2+ = data  (leave blank = skip that field for that lesson)

    Returns: {content_key: [(lesson_name, new_value), ...]}
    """
    headers = [c.value for c in ws[1]]
    if not headers or headers[0] is None:
        return {}

    updates: dict[str, list[tuple[str, str]]] = OrderedDict()
    for row in ws.iter_rows(min_row=2, values_only=True):
        lesson_name = row[0]
        if lesson_name is None or str(lesson_name).strip() == "":
            continue
        lesson_name = str(lesson_name).strip()
        if lesson_name.startswith("("):  # skip example / placeholder rows
            continue

        for col_idx in range(1, len(headers)):
            header = headers[col_idx] if col_idx < len(headers) else None
            val    = row[col_idx]     if col_idx < len(row)     else None
            if header is None or val is None or str(val).strip() == "":
                continue
            key = str(header).strip()
            updates.setdefault(key, []).append((lesson_name, str(val).strip()))

    return updates


def build_cte_pair(
    sheet_name: str,
    content_key: str,
    json_index: int,
    rows: list[tuple[str, str]],
    prefix: str,
) -> tuple[str, list[str]]:
    """
    Build one pair of UPDATE CTEs for a single content_key.

    Returns (cte_sql_fragment, [select_count_expressions])
    """
    data_cte = f"{prefix}_data"
    lc_cte   = f"{prefix}_lc"
    lv_cte   = f"{prefix}_lv"

    values_lines = [f"        ('{sq(name)}', '{sq(val)}')" for name, val in rows]
    values_str   = ",\n".join(values_lines)

    cte = (
        f"-- {sheet_name}: update '{content_key}' ({len(rows)} lesson(s))\n"
        f"{data_cte} AS (\n"
        f"    SELECT * FROM (VALUES\n"
        f"{values_str}\n"
        f"    ) AS v(lesson_name, new_value)\n"
        f"),\n"
        f"{lc_cte} AS (\n"
        f"    UPDATE bantrly.lesson_content lc\n"
        f"    SET content_value = ud.new_value,\n"
        f"        updated_at = NOW()\n"
        f"    FROM bantrly.lesson l\n"
        f"    JOIN {data_cte} ud ON ud.lesson_name = l.lesson_name\n"
        f"    WHERE lc.lesson_id = l.lesson_id\n"
        f"      AND lc.content_key = '{content_key}'\n"
        f"      AND l.difficulty_level = '{DIFFICULTY}'\n"
        f"      AND l.created_by = '{CREATED_BY}'\n"
        f"      AND l.created_by_role = '{ROLE}'\n"
        f"    RETURNING l.lesson_name\n"
        f"),\n"
        f"{lv_cte} AS (\n"
        f"    UPDATE bantrly.lesson_version_new lvn\n"
        f"    SET content = jsonb_set(\n"
        f"            lvn.content::jsonb,\n"
        f"            '{{contents,{json_index},content_value}}',\n"
        f"            to_jsonb(ud.new_value)\n"
        f"        )\n"
        f"    FROM bantrly.lesson l\n"
        f"    JOIN {data_cte} ud ON ud.lesson_name = l.lesson_name\n"
        f"    WHERE lvn.lesson_id = l.lesson_id\n"
        f"      AND lvn.is_latest = true\n"
        f"      AND l.difficulty_level = '{DIFFICULTY}'\n"
        f"      AND l.created_by = '{CREATED_BY}'\n"
        f"      AND l.created_by_role = '{ROLE}'\n"
        f"    RETURNING lvn.title\n"
        f")"
    )

    selects = [
        f"    (SELECT COUNT(*) FROM {lc_cte}) AS {prefix}_content_updated",
        f"    (SELECT COUNT(*) FROM {lv_cte}) AS {prefix}_version_updated",
    ]
    return cte, selects


def generate_sql(excel_path: str) -> str | None:
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    all_ctes: list[str]    = []
    all_selects: list[str] = []
    total_fields = 0

    for sheet_name, key_index_pairs in MODULE_KEYS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        updates = read_sheet(ws)
        if not updates:
            continue

        key_to_idx = {k: i for k, i in key_index_pairs}
        # short sheet code: "M1 - Reading Paragraph" → "m1"
        sheet_code = sheet_name.split(" - ")[0].strip().lower().replace(" ", "")

        for content_key, rows in updates.items():
            if content_key not in key_to_idx:
                print(f"  WARNING: '{content_key}' not a known key for {sheet_name} — skipped")
                continue
            json_idx = key_to_idx[content_key]
            key_code = content_key.replace("_", "")
            prefix   = f"{sheet_code}_{key_code}"

            cte_sql, select_cols = build_cte_pair(sheet_name, content_key, json_idx, rows, prefix)
            all_ctes.append(cte_sql)
            all_selects.extend(select_cols)
            total_fields += 1
            print(f"  + {sheet_name} / '{content_key}'  ({len(rows)} row(s))")

    if not all_ctes:
        print("No updates found in the Excel file.")
        return None

    ctes_block    = ",\n\n".join(all_ctes)
    selects_block = ",\n".join(all_selects)

    sql = (
        f"-- Generated by generate_update_sql.py\n"
        f"-- Filters: difficulty={DIFFICULTY} | created_by={CREATED_BY} | role={ROLE}\n"
        f"-- Fields updated: {total_fields}\n"
        f"-- Only the specified content_keys are changed; all other values remain untouched.\n\n"
        f"BEGIN;\n\n"
        f"WITH\n"
        f"{ctes_block}\n\n"
        f"SELECT\n"
        f"{selects_block};\n\n"
        f"COMMIT;\n"
    )
    return sql


# ─── Template creator ──────────────────────────────────────────────────────────
EXAMPLE_VALUES = {
    "passage":               "New passage text here...",
    "example":               "New example text here...",
    "example_question":      "New example question?",
    "example_content":       "New example content...",
    "keywords":              "word1, word2, word3",
    "theme":                 "New theme topic...",
    "rapid_questions":       "Question 1?.Question 2?.Question 3?",
    "example_heading":       "New example heading...",
    "topic_overview":        "New topic overview...",
    "preparation_guidelines":"Guideline 1.Guideline 2.",
    "tongue_twister":        "New tongue twister text...",
    "assigned_text":         "New poem text here...",
    "topic":                 "New debate topic statement.",
    "user_side":             "Side the student argues for.",
    "ai_side":               "Side the AI argues for.",
    "content":               "Debate prompt / conversation starter.",
    "example_topic":         "Ice Cream is better than Pizza",
    "example_user_side":     "Ice cream is cold and sweet...",
    "example_ai_side":       "Pizza has cheese and toppings...",
    "instructions":          "New instructions text...",
    "example_instruction":   "New example instruction...",
}

def create_template(out_path: str = "update_template.xlsx") -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill("solid", fgColor="2E5090")
    note_font = Font(italic=True, color="999999")
    note_fill = PatternFill("solid", fgColor="FFF2CC")
    center    = Alignment(horizontal="center", wrap_text=True)

    for sheet_name, key_index_pairs in MODULE_KEYS.items():
        ws = wb.create_sheet(sheet_name)
        headers = ["lesson_name"] + [k for k, _ in key_index_pairs]

        # Header row
        for col, header in enumerate(headers, 1):
            cell       = ws.cell(row=1, column=col, value=header)
            cell.font  = hdr_font
            cell.fill  = hdr_fill
            cell.alignment = center
            ws.column_dimensions[get_column_letter(col)].width = max(22, len(header) + 6)

        # Example row (greyed out — not real data)
        ws.cell(row=2, column=1, value="(example lesson name — delete this row)").font = note_font
        ws.row_dimensions[2].fill = note_fill  # type: ignore
        for col, (key, _) in enumerate(key_index_pairs, 2):
            cell       = ws.cell(row=2, column=col, value=EXAMPLE_VALUES.get(key, f"(example {key})"))
            cell.font  = note_font
            cell.fill  = note_fill
            cell.alignment = Alignment(wrap_text=True)

        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 40
        ws.freeze_panes = "A3"

    p = Path(out_path)
    wb.save(p)
    print(f"Template saved: {p}")
    print()
    print("How to use:")
    print("  1. Fill in 'lesson_name' (must match exactly what's in the database).")
    print("  2. Fill in ONLY the columns you want to update.")
    print("  3. Leave cells blank to keep existing values unchanged.")
    print("  4. You can use multiple sheets in one file.")
    print("  5. Delete or ignore the yellow example row.")


# ─── Main ─────────────────────────────────────────────────────────────────────
def main() -> None:
    args = sys.argv[1:]

    if "--create-template" in args:
        out = next((a for a in args if not a.startswith("--")), "update_template.xlsx")
        create_template(out)
        return

    excel_file = next((a for a in args if not a.startswith("--")), DEFAULT_EXCEL)
    p = Path(excel_file)
    if not p.exists():
        sys.exit(
            f"'{excel_file}' not found.\n"
            "Run with --create-template to generate a blank template."
        )

    print(f"Reading:  {p}")
    sql = generate_sql(str(p))

    if sql:
        out = Path(OUTPUT_FILE)
        out.write_text(sql, encoding="utf-8")
        print(f"\nSaved:  {out}  ({len(sql):,} chars)")

    print("\nDone.")


if __name__ == "__main__":
    main()
