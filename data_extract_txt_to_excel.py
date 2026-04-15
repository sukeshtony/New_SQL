import os
import re
import json
import random
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from google import genai
from google.genai import types


DOC_PATH = "source_lessons.txt"
EXCEL_TEMPLATE_PATH = "lesson_template.xlsx"
OUTPUT_PATH = "lesson_template_filled.xlsx"

MODEL_NAME = "gemini-2.5-pro"
THINKING_BUDGET = 32768
TEMPERATURE = 0.1

IMAGE_URL_TEMPLATE = (
    "https://ai-voice-app-audio-storage.s3.us-east-1.amazonaws.com/lesson_images/{n}.png"
)

client = genai.Client()


def read_text_file(path: str) -> str:
    return Path(path).read_text(encoding="utf-8")


def normalize_header(text: str) -> str:
    text = (text or "").strip().lower()
    text = re.sub(r"[\[\]\(\)\-_/|]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text


def random_image_url() -> str:
    return IMAGE_URL_TEMPLATE.format(n=random.randint(1, 100))


def split_modules(full_text: str) -> List[Dict[str, str]]:
    pattern = re.compile(r"(?=Module\s+\d+\s*:)", re.IGNORECASE)
    parts = [p.strip() for p in pattern.split(full_text) if p.strip()]

    modules = []
    for part in parts:
        first_line = part.splitlines()[0].strip()
        m = re.match(r"Module\s+(\d+)\s*:\s*(.+)", first_line, re.IGNORECASE)
        if not m:
            continue
        modules.append(
            {
                "module_number": int(m.group(1)),
                "module_title": m.group(2).strip(),
                "text": part,
            }
        )
    return modules


def get_sheet_by_module_number(workbook, module_number: int):
    idx = module_number - 1
    if idx < 0 or idx >= len(workbook.worksheets):
        raise IndexError(f"No sheet found for module {module_number}")
    return workbook.worksheets[idx]


def build_json_schema_for_module(module_number: int) -> Dict[str, Any]:
    common_lesson = {
        "type": "object",
        "properties": {
            "lesson_number": {"type": "integer"},
            "lesson_title": {"type": "string"},
            "standards_alignment": {"type": "string"},
            "lesson_image_url": {"type": "string"},
        },
        "required": ["lesson_number", "lesson_title", "standards_alignment"],
    }

    if module_number == 1:
        common_lesson["properties"].update(
            {
                "passage": {"type": "string"},
            }
        )
        common_lesson["required"].extend(["passage"])

    elif module_number == 2:
        common_lesson["properties"].update(
            {
                "keywords": {
                    "type": "array",
                    "items": {"type": "string"},
                    "minItems": 5,
                    "maxItems": 5,
                }
            }
        )
        common_lesson["required"].extend(["keywords"])

    elif module_number == 3:
        common_lesson["properties"].update(
            {
                "theme": {"type": "string"},
                "questions": {
                    "type": "array",
                    "items": {"type": "string"},
                    "minItems": 5,
                    "maxItems": 5,
                },
            }
        )
        common_lesson["required"].extend(["theme", "questions"])

    elif module_number == 4:
        common_lesson["properties"].update(
            {
                "topic_overview": {"type": "string"},
                "preparation_guidelines": {
                    "type": "array",
                    "items": {"type": "string"},
                },
            }
        )
        common_lesson["required"].extend(["topic_overview", "preparation_guidelines"])

    elif module_number == 5:
        common_lesson["properties"].update(
            {
                "tongue_twister_lines": {
                    "type": "array",
                    "items": {"type": "string"},
                    "minItems": 3,
                },
            }
        )
        common_lesson["required"].append("tongue_twister_lines")

    elif module_number == 6:
        common_lesson["properties"].update(
            {
                "theme": {"type": "string"},
                "poem_lines": {
                    "type": "array",
                    "items": {"type": "string"},
                    "minItems": 1,
                },
            }
        )
        common_lesson["required"].extend(["theme", "poem_lines"])

    elif module_number == 7:
        common_lesson["properties"].update(
            {
                "debate_statement": {"type": "string"},
                "ai_dialogue_prompts": {
                    "type": "array",
                    "items": {"type": "string"},
                    "minItems": 3,
                },
            }
        )
        common_lesson["required"].extend(["debate_statement", "ai_dialogue_prompts"])

    elif module_number == 8:
        common_lesson["properties"].update(
            {
                "overview": {"type": "string"},
                "instructions": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "sample": {"type": "string"},
            }
        )
        common_lesson["required"].extend(["overview", "instructions", "sample"])

    else:
        raise ValueError(f"Unsupported module number: {module_number}")

    return {
        "type": "object",
        "properties": {
            "module_number": {"type": "integer"},
            "module_title": {"type": "string"},
            "lessons": {
                "type": "array",
                "items": common_lesson,
            },
        },
        "required": ["module_number", "module_title", "lessons"],
    }


def build_prompt_for_module(module_number: int, module_text: str) -> str:
    base = f"""
You are extracting lesson content from one module of a curriculum document.

Return ONLY valid JSON that matches the provided response schema.

Critical rules:
- Extract every lesson in this module.
- Keep original lesson order.
- Do not skip lessons.
- Do not rewrite or summarize unless needed only to fit the schema.
- Preserve original punctuation and capitalization.
- For fields that need multiple points, return them as arrays in the original order.
- For Silly Debates, each AI dialogue prompt must remain a separate array item.
- For poems, each poem line must be a separate array item.
- For tongue twisters, each repeated line must be a separate array item.
- Do not invent missing content.
- Do not extract SEL Focus at all.
- For Module 2, do not extract example content at all because the Excel already contains the same fixed example.
- Set lesson_image_url to an empty string in the JSON. Python will fill it later.

Module-specific notes:
"""

    if module_number == 1:
        base += """
Module 1:
- Extract lesson_number, lesson_title, passage, standards_alignment.
"""
    elif module_number == 2:
        base += """
Module 2:
- Extract lesson_number, lesson_title, keywords (5 items), standards_alignment.
- Do not include example content.
"""
    elif module_number == 3:
        base += """
Module 3:
- Extract lesson_number, lesson_title, theme, questions (5 items), standards_alignment.
- Do not include SEL Focus.
"""
    elif module_number == 4:
        base += """
Module 4:
- Extract lesson_number, lesson_title, topic_overview, preparation_guidelines, standards_alignment.
"""
    elif module_number == 5:
        base += """
Module 5:
- Extract lesson_number, lesson_title, tongue_twister_lines, standards_alignment.
"""
    elif module_number == 6:
        base += """
Module 6:
- Extract lesson_number, lesson_title, theme, poem_lines, standards_alignment.
"""
    elif module_number == 7:
        base += """
Module 7:
- Extract lesson_number, lesson_title, debate_statement, ai_dialogue_prompts, standards_alignment.
- Keep every debate prompt as a separate item.
"""
    elif module_number == 8:
        base += """
Module 8:
- Extract lesson_number, lesson_title, overview, instructions, sample, standards_alignment.
- Keep instructions as separate items when possible.
"""

    base += f"""

Module text:
{module_text}
"""
    return base.strip()


def call_gemini_extract(module_number: int, module_title: str, module_text: str) -> Dict[str, Any]:
    schema = build_json_schema_for_module(module_number)
    prompt = build_prompt_for_module(module_number, module_text)

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=TEMPERATURE,
            response_mime_type="application/json",
            response_schema=schema,
            thinking_config=types.ThinkingConfig(
                thinking_budget=THINKING_BUDGET
            ),
        ),
    )

    data = json.loads(response.text)

    for lesson in data.get("lessons", []):
        lesson["lesson_image_url"] = random_image_url()

    return data


def join_lines(value: Optional[List[str]]) -> str:
    if not value:
        return ""
    return "\n".join(v.strip() for v in value if str(v).strip())


def lesson_to_flat_row(module_number: int, lesson: Dict[str, Any]) -> Dict[str, str]:
    row = {
        "lesson_number": str(lesson.get("lesson_number", "")),
        "lesson_title": lesson.get("lesson_title", ""),
        "standards_alignment": lesson.get("standards_alignment", ""),
        "lesson_image_url": lesson.get("lesson_image_url", ""),
    }

    if module_number == 1:
        row.update({
            "passage": lesson.get("passage", ""),
        })

    elif module_number == 2:
        row.update({
            "keywords": ", ".join(lesson.get("keywords", [])),
        })

    elif module_number == 3:
        row.update({
            "theme": lesson.get("theme", ""),
            "questions": join_lines(lesson.get("questions", [])),
        })

    elif module_number == 4:
        row.update({
            "topic_overview": lesson.get("topic_overview", ""),
            "preparation_guidelines": join_lines(lesson.get("preparation_guidelines", [])),
        })

    elif module_number == 5:
        row.update({
            "tongue_twister_lines": join_lines(lesson.get("tongue_twister_lines", [])),
        })

    elif module_number == 6:
        row.update({
            "theme": lesson.get("theme", ""),
            "poem_lines": join_lines(lesson.get("poem_lines", [])),
        })

    elif module_number == 7:
        row.update({
            "debate_statement": lesson.get("debate_statement", ""),
            "ai_dialogue_prompts": join_lines(lesson.get("ai_dialogue_prompts", [])),
        })

    elif module_number == 8:
        row.update({
            "overview": lesson.get("overview", ""),
            "instructions": join_lines(lesson.get("instructions", [])),
            "sample": lesson.get("sample", ""),
        })

    return row


def build_header_map(sheet) -> Dict[str, int]:
    header_map = {}
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=col).value
        if value:
            header_map[normalize_header(str(value))] = col
    return header_map


def find_target_column(header_map: Dict[str, int], candidates: List[str]) -> Optional[int]:
    for candidate in candidates:
        key = normalize_header(candidate)
        if key in header_map:
            return header_map[key]
    return None


FIELD_CANDIDATES = {
    "lesson_number": ["lesson number", "lesson_no", "lesson no", "number"],
    "lesson_title": ["lesson title", "title", "lesson name", "lesson"],
    "passage": ["passage", "content", "lesson content", "story", "paragraph"],
    "standards_alignment": ["standards alignment", "standards", "alignment"],
    "lesson_image_url": ["lesson_image_url", "lesson image url", "image url"],
    "keywords": ["keywords", "keyword"],
    "theme": ["theme"],
    "questions": ["questions", "qa", "q&a", "question set"],
    "topic_overview": ["topic overview", "overview"],
    "preparation_guidelines": ["preparation guidelines", "guidelines", "prep guidelines"],
    "tongue_twister_lines": ["tongue twister", "tongue twister lines", "lines"],
    "poem_lines": ["poem", "poem lines", "lines"],
    "debate_statement": ["debate statement", "statement"],
    "ai_dialogue_prompts": ["ai dialogue prompts", "prompts", "dialogue prompts"],
    "instructions": ["instructions"],
    "sample": ["sample"],
}


def write_lessons_to_sheet(sheet, module_number: int, lessons: List[Dict[str, Any]]) -> None:
    header_map = build_header_map(sheet)
    start_row = sheet.max_row + 1

    for i, lesson in enumerate(lessons, start=start_row):
        flat = lesson_to_flat_row(module_number, lesson)

        for field_name, value in flat.items():
            col = find_target_column(header_map, FIELD_CANDIDATES.get(field_name, [field_name]))
            if col is not None:
                sheet.cell(row=i, column=col).value = value


def validate_module_counts(extracted_modules: List[Dict[str, Any]]) -> None:
    for mod in extracted_modules:
        if not mod.get("lessons"):
            raise ValueError(f"Module {mod.get('module_number')} has no extracted lessons.")


def main():
    source_text = read_text_file(DOC_PATH)
    modules = split_modules(source_text)

    if len(modules) != 8:
        print(f"Warning: found {len(modules)} modules, expected 8.")

    wb = load_workbook(EXCEL_TEMPLATE_PATH)

    extracted_modules = []
    for module in modules:
        module_number = module["module_number"]
        module_title = module["module_title"]
        module_text = module["text"]

        print(f"Extracting Module {module_number}: {module_title}")
        data = call_gemini_extract(module_number, module_title, module_text)
        extracted_modules.append(data)

    validate_module_counts(extracted_modules)

    for mod in extracted_modules:
        sheet = get_sheet_by_module_number(wb, mod["module_number"])
        write_lessons_to_sheet(sheet, mod["module_number"], mod["lessons"])

    wb.save(OUTPUT_PATH)
    print(f"Done. Saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()