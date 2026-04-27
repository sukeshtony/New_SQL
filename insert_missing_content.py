"""
insert_missing_content.py
────────────────────────────────────────────────────────────────
Generates SQL to INSERT missing lesson_content rows and APPEND
missing items to the lesson_version_new JSON contents array.

Use this when content keys were forgotten during the initial insert.
(generate_update_sql.py only works when rows already exist — this
 script handles the case where they are completely absent.)

Each generated query is safe to re-run: INSERT is guarded by
WHERE NOT EXISTS, and the JSON append checks the array first.

Usage:
    python insert_missing_content.py                     # default Excel
    python insert_missing_content.py my_file.xlsx        # custom file
    python insert_missing_content.py my_file.xlsx 1 4    # modules 1 and 4
    python insert_missing_content.py --list              # show modules
    python insert_missing_content.py --create-template   # blank template

Output: insert_missing_output.sql
"""

import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl not installed.  Run: pip install openpyxl")

CONFIG_SHEET  = "Config"
DEFAULT_EXCEL = "insert_missing_template.xlsx"
OUTPUT_FILE   = "insert_missing_output.sql"

# ─── Module definitions ────────────────────────────────────────────────────────
# sheet_name → ordered list of (content_key, json_array_index)
MODULE_KEYS: dict[str, list[tuple[str, int]]] = {
    "M1 - Reading Paragraph": [
        ("passage",      0),
        ("example",      1),
    ],
    "M2 - Story Building": [
        ("example_question", 0),
        ("example_content",  1),
        ("keywords",         2),
    ],
    "M3 - QA Bursts": [
        ("theme",           0),
        ("rapid_questions", 1),
        ("example_heading", 2),
        ("example_content", 3),
    ],
    "M4 - Presentation": [
        ("topic_overview",         0),
        ("preparation_guidelines", 1),
        ("example_heading",        2),
        ("example_content",        3),
    ],
    "M5 - Tongue Twisters": [
        ("example",        0),
        ("tongue_twister", 1),
    ],
    "M6 - Poems": [
        ("example",       0),
        ("assigned_text", 1),
    ],
    "M7 - Silly Debate": [
        ("topic",             0),
        ("user_side",         1),
        ("ai_side",           2),
        ("content",           3),
        ("example_topic",     4),
        ("example_user_side", 5),
        ("example_ai_side",   6),
    ],
    "M8 - Audio Postcards": [
        ("topic",               0),
        ("instructions",        1),
        ("example_topic",       2),
        ("example_instruction", 3),
        ("example_content",     4),
    ],
}

VALID_MODULES = {int(s.split(" - ")[0][1:]): s for s in MODULE_KEYS}


# ─── Helpers ───────────────────────────────────────────────────────────────────
def sq(val: str) -> str:
    return str(val).replace("'", "''")


def read_config(wb) -> dict[str, str]:
    if CONFIG_SHEET not in wb.sheetnames:
        sys.exit(
            f"'{CONFIG_SHEET}' sheet not found.\n"
            "Run with --create-template to regenerate the template."
        )
    ws  = wb[CONFIG_SHEET]
    cfg: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key, val = row[0], row[1]
        if key and val is not None and str(key).strip():
            cfg[str(key).strip().lower()] = str(val).strip()
    missing = [f for f in ("created_by", "role", "difficulty") if f not in cfg]
    if missing:
        sys.exit(f"Config sheet is missing: {', '.join(missing)}")
    return cfg


def read_sheet_wide(ws, all_keys: list[str]) -> list[dict]:
    """
    Read a module sheet in wide format.
    Row 1 = headers (lesson_name, key1, key2 ...)
    Row 2+ = data   (blank cell = this key is NOT missing for this lesson)

    Returns list of dicts: [{lesson_name, key1: val_or_None, ...}]
    Only includes rows that have at least one non-blank key value.
    """
    headers = [c.value for c in ws[1]]
    if not headers or headers[0] is None:
        return []

    rows_out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        lesson_name = row[0]
        if lesson_name is None or str(lesson_name).strip() == "":
            continue
        lesson_name = str(lesson_name).strip()
        if lesson_name.startswith("("):
            continue

        record: dict = {"lesson_name": lesson_name}
        has_any = False
        for key in all_keys:
            if key in headers:
                col_idx = headers.index(key)
                val = row[col_idx] if col_idx < len(row) else None
                val = str(val).strip() if val is not None and str(val).strip() else None
                record[key] = val
                if val:
                    has_any = True
            else:
                record[key] = None

        if has_any:
            rows_out.append(record)
    return rows_out


# ─── SQL builders ──────────────────────────────────────────────────────────────
def build_module_block(
    sheet_name: str,
    key_index_pairs: list[tuple[str, int]],
    lessons: list[dict],
    cfg: dict[str, str],
) -> tuple[str, list[str]]:
    """
    Build the full CTE block for one module:
      - one wide data CTE
      - one INSERT CTE per key (with NOT EXISTS guard)
      - one version UPDATE CTE (appends all missing keys in index order)

    Returns (cte_sql, [select_count_expressions])
    """
    created_by = cfg["created_by"]
    role       = cfg["role"]
    difficulty = cfg["difficulty"]

    mod_code  = sheet_name.split(" - ")[0].strip().lower()   # m1, m2 …
    all_keys  = [k for k, _ in key_index_pairs]

    # ── Wide data CTE ──────────────────────────────────────────────────────────
    # Columns: lesson_name, key1, key2, ...
    # NULL sentinel for missing/blank values
    col_names = ", ".join(["lesson_name"] + all_keys)
    value_rows = []
    for les in lessons:
        cols = [f"'{sq(les['lesson_name'])}'"]
        for key in all_keys:
            v = les.get(key)
            cols.append(f"'{sq(v)}'" if v else "NULL")
        value_rows.append("        (" + ", ".join(cols) + ")")
    values_str = ",\n".join(value_rows)

    data_cte_name = f"{mod_code}_data"
    data_cte = (
        f"-- {sheet_name}: insert missing content ({len(lessons)} lesson(s))\n"
        f"{data_cte_name} AS (\n"
        f"    SELECT * FROM (VALUES\n"
        f"{values_str}\n"
        f"    ) AS v({col_names})\n"
        f")"
    )

    all_ctes  = [data_cte]
    all_selects = []

    # ── INSERT CTEs (one per key) ──────────────────────────────────────────────
    insert_cte_names = []
    for key, _ in key_index_pairs:
        # Only generate if at least one lesson has a value for this key
        if not any(les.get(key) for les in lessons):
            continue

        ins_name = f"{mod_code}_{key.replace('_', '')}_ins"
        insert_cte_names.append(ins_name)

        ins_cte = (
            f"{ins_name} AS (\n"
            f"    INSERT INTO bantrly.lesson_content\n"
            f"        (content_id, lesson_id, content_key, content_value, created_at, updated_at)\n"
            f"    SELECT\n"
            f"        gen_random_uuid(), l.lesson_id,\n"
            f"        '{key}', d.{key},\n"
            f"        NOW(), NOW()\n"
            f"    FROM bantrly.lesson l\n"
            f"    JOIN {data_cte_name} d ON d.lesson_name = l.lesson_name\n"
            f"    WHERE l.difficulty_level = '{difficulty}'\n"
            f"      AND l.created_by       = '{created_by}'\n"
            f"      AND l.created_by_role  = '{role}'\n"
            f"      AND d.{key} IS NOT NULL\n"
            f"      AND NOT EXISTS (\n"
            f"          SELECT 1 FROM bantrly.lesson_content lc\n"
            f"          WHERE lc.lesson_id   = l.lesson_id\n"
            f"            AND lc.content_key = '{key}'\n"
            f"      )\n"
            f"    RETURNING lesson_id\n"
            f")"
        )
        all_ctes.append(ins_cte)
        all_selects.append(
            f"    (SELECT COUNT(*) FROM {ins_name}) AS {mod_code}_{key.replace('_','')}_inserted"
        )

    # ── Version UPDATE CTE (one per module, appends all keys) ─────────────────
    ver_name = f"{mod_code}_version_upd"

    # Build the append expression: one CASE block per key, in index order
    append_parts = []
    for key, _ in key_index_pairs:
        if not any(les.get(key) for les in lessons):
            continue
        append_parts.append(
            f"        || CASE\n"
            f"             WHEN d.{key} IS NOT NULL\n"
            f"              AND NOT (lvn.content->'contents' @> "
            f"'[{{\"content_key\":\"{key}\"}}]'::jsonb)\n"
            f"             THEN jsonb_build_array(\n"
            f"                      jsonb_build_object(\n"
            f"                          'content_key',   '{key}',\n"
            f"                          'content_value', d.{key}\n"
            f"                      )\n"
            f"                  )\n"
            f"             ELSE '[]'::jsonb\n"
            f"           END"
        )

    if append_parts:
        append_expr = "\n".join(append_parts)
        ver_cte = (
            f"{ver_name} AS (\n"
            f"    UPDATE bantrly.lesson_version_new lvn\n"
            f"    SET content = jsonb_set(\n"
            f"            lvn.content::jsonb,\n"
            f"            '{{contents}}',\n"
            f"            (lvn.content->'contents')\n"
            f"{append_expr}\n"
            f"        )\n"
            f"    FROM bantrly.lesson l\n"
            f"    JOIN {data_cte_name} d ON d.lesson_name = l.lesson_name\n"
            f"    WHERE lvn.lesson_id    = l.lesson_id\n"
            f"      AND lvn.is_latest   = true\n"
            f"      AND l.difficulty_level = '{difficulty}'\n"
            f"      AND l.created_by       = '{created_by}'\n"
            f"      AND l.created_by_role  = '{role}'\n"
            f"    RETURNING lvn.title\n"
            f")"
        )
        all_ctes.append(ver_cte)
        all_selects.append(
            f"    (SELECT COUNT(*) FROM {ver_name}) AS {mod_code}_versions_updated"
        )

    return ",\n\n".join(all_ctes), all_selects


# ─── Main generator ───────────────────────────────────────────────────────────
def generate_sql(excel_path: str, selected: set[int] | None = None) -> str | None:
    wb  = openpyxl.load_workbook(excel_path, data_only=True)
    cfg = read_config(wb)

    all_ctes:    list[str] = []
    all_selects: list[str] = []

    for sheet_name, key_index_pairs in MODULE_KEYS.items():
        mod_num = int(sheet_name.split(" - ")[0][1:])
        if selected and mod_num not in selected:
            continue
        if sheet_name not in wb.sheetnames:
            continue

        all_keys = [k for k, _ in key_index_pairs]
        lessons  = read_sheet_wide(wb[sheet_name], all_keys)
        if not lessons:
            continue

        keys_present = [k for k in all_keys if any(l.get(k) for l in lessons)]
        print(f"  + {sheet_name}  ({len(lessons)} lesson(s), keys: {keys_present})")

        cte_block, selects = build_module_block(sheet_name, key_index_pairs, lessons, cfg)
        all_ctes.append(cte_block)
        all_selects.extend(selects)

    if not all_ctes:
        print("No data found in the Excel file.")
        return None

    sql = (
        f"-- Generated by insert_missing_content.py\n"
        f"-- Filters: difficulty={cfg['difficulty']} | created_by={cfg['created_by']} | role={cfg['role']}\n"
        f"-- Inserts missing lesson_content rows and appends to lesson_version_new JSON.\n"
        f"-- Safe to re-run: INSERT uses NOT EXISTS guard, JSON append checks array first.\n\n"
        f"BEGIN;\n\n"
        f"WITH\n"
        + ",\n\n".join(all_ctes)
        + "\n\nSELECT\n"
        + ",\n".join(all_selects)
        + ";\n\nCOMMIT;\n"
    )
    return sql


# ─── Template creator ──────────────────────────────────────────────────────────
EXAMPLE_VALUES = {
    "passage":               "Full paragraph text here...",
    "example":               "Read the following paragraph clearly and with appropriate expression.",
    "example_question":      "Example question here?",
    "example_content":       "Example answer content here...",
    "keywords":              "word1, word2, word3",
    "theme":                 "Theme topic here...",
    "rapid_questions":       "Question 1?.Question 2?.Question 3?",
    "example_heading":       "Hey, what is your favorite...?",
    "topic_overview":        "Topic overview text here...",
    "preparation_guidelines":"Guideline 1.Guideline 2.",
    "tongue_twister":        "Tongue twister text here...",
    "assigned_text":         "Poem text here...",
    "topic":                 "Debate topic statement.",
    "user_side":             "Side the student argues for.",
    "ai_side":               "Side the AI argues for.",
    "content":               "Debate prompt / conversation starter.",
    "example_topic":         "Example topic here...",
    "example_user_side":     "Example user side argument...",
    "example_ai_side":       "Example AI side argument...",
    "instructions":          "Instructions text here...",
    "example_instruction":   "Example instruction here...",
}

def create_template(out_path: str = DEFAULT_EXCEL) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill("solid", fgColor="1F6B3A")   # green — distinct from update template
    note_font = Font(italic=True, color="999999")
    note_fill = PatternFill("solid", fgColor="E2EFDA")
    center    = Alignment(horizontal="center", wrap_text=True)

    # ── Config sheet ────────────────────────────────────────────────────────────
    ws_cfg = wb.create_sheet(CONFIG_SHEET, 0)
    ws_cfg.column_dimensions["A"].width = 22
    ws_cfg.column_dimensions["B"].width = 45
    ws_cfg.column_dimensions["C"].width = 44

    for col, label in enumerate(("Field", "Value", "Description"), 1):
        cell      = ws_cfg.cell(row=1, column=col, value=label)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center

    cfg_rows = [
        ("created_by", "ee42009d-d83e-4c12-abd1-2d8fff809b18", "UUID of the teacher who created the lessons"),
        ("role",       "TEACHER",                               "Role filter  (e.g. TEACHER)"),
        ("difficulty", "EASY",                                  "Difficulty level  (e.g. EASY, MEDIUM, HARD)"),
    ]
    for row_idx, (field, default_val, desc) in enumerate(cfg_rows, 2):
        ws_cfg.cell(row=row_idx, column=1, value=field).font = Font(bold=True)
        ws_cfg.cell(row=row_idx, column=2, value=default_val)
        ws_cfg.cell(row=row_idx, column=3, value=desc).font  = Font(italic=True, color="666666")

    ws_cfg.freeze_panes = "B2"

    # ── Module sheets ───────────────────────────────────────────────────────────
    for sheet_name, key_index_pairs in MODULE_KEYS.items():
        ws = wb.create_sheet(sheet_name)
        headers = ["lesson_name"] + [k for k, _ in key_index_pairs]

        for col, header in enumerate(headers, 1):
            cell           = ws.cell(row=1, column=col, value=header)
            cell.font      = hdr_font
            cell.fill      = hdr_fill
            cell.alignment = center
            ws.column_dimensions[get_column_letter(col)].width = max(24, len(header) + 6)

        # Example row
        ex_lesson = ws.cell(row=2, column=1, value="(example — delete this row)")
        ex_lesson.font = note_font
        ex_lesson.fill = note_fill
        for col, (key, _) in enumerate(key_index_pairs, 2):
            cell       = ws.cell(row=2, column=col, value=EXAMPLE_VALUES.get(key, f"(example {key})"))
            cell.font  = note_font
            cell.fill  = note_fill
            cell.alignment = Alignment(wrap_text=True)

        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 35
        ws.freeze_panes = "A3"

    wb.save(Path(out_path))
    print(f"Template saved: {out_path}")
    print()
    print("How to use:")
    print("  1. Fill in Config sheet (created_by, role, difficulty).")
    print("  2. In each module sheet, fill in lesson_name and ONLY the keys that are missing.")
    print("  3. Leave a cell blank = that key already exists for that lesson (will be skipped).")
    print("  4. Delete or ignore the green example row.")


# ─── Main ─────────────────────────────────────────────────────────────────────
def main() -> None:
    args = sys.argv[1:]

    if "--create-template" in args:
        positional = [a for a in args if not a.startswith("--")]
        out = positional[0] if positional else DEFAULT_EXCEL
        create_template(out)
        return

    if "--list" in args:
        print("Available modules:")
        for num, sheet in VALID_MODULES.items():
            print(f"  {num}  {sheet}")
        return

    positional = [a for a in args if not a.startswith("--")]
    mod_args   = [a for a in positional if a.isdigit()]
    file_args  = [a for a in positional if not a.isdigit()]

    excel_file = file_args[0] if file_args else DEFAULT_EXCEL

    selected: set[int] | None = None
    if mod_args:
        selected = set()
        for m in mod_args:
            n = int(m)
            if n not in VALID_MODULES:
                sys.exit(f"Module {n} does not exist. Use --list to see valid modules.")
            selected.add(n)

    p = Path(excel_file)
    if not p.exists():
        sys.exit(
            f"'{excel_file}' not found.\n"
            "Run with --create-template to generate a blank template."
        )

    label = f"modules {sorted(selected)}" if selected else "all modules"
    print(f"Reading:  {p}  [{label}]")
    sql = generate_sql(str(p), selected)

    if sql:
        out = Path(OUTPUT_FILE)
        out.write_text(sql, encoding="utf-8")
        print(f"\nSaved:  {out}  ({len(sql):,} chars)")

    print("\nDone.")


if __name__ == "__main__":
    main()
